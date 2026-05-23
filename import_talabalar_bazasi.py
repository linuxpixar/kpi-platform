"""
Talabalar bazasini import qilish skripti.
Fayl: Talabalar-21_05_2026_02_12_01.xlsx

Ishlatish:
  1. Excel faylni shu papkaga ko'chiring
  2. python3 import_talabalar_bazasi.py

Ustunlar: F.I.O. | Kurs | Fakultet | Guruh
"""
import os, sys, django
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'sertifikat_platform.settings')
django.setup()

from openpyxl import load_workbook
from core.models import Fakultet, Guruh, Talaba

# ── FAYL JOYLASHUVI ───────────────────────────────────────────────
FAYL = os.path.join(os.path.dirname(__file__), 'Talabalar-21_05_2026_02_12_01.xlsx')

if not os.path.exists(FAYL):
    print(f"❌ Fayl topilmadi: {FAYL}")
    print("   Excel faylni shu papkaga ko'chiring:")
    print(f"   {os.path.dirname(__file__)}")
    sys.exit(1)

# ── KURS RAQAMI ───────────────────────────────────────────────────
def kurs_raqam(kurs_str):
    mapping = {
        '1-kurs': 1, '2-kurs': 2, '3-kurs': 3,
        '4-kurs': 4, 'magistratura': 5,
        '1': 1, '2': 2, '3': 3, '4': 4, '5': 5
    }
    if not kurs_str:
        return 1
    return mapping.get(str(kurs_str).strip().lower(), 1)

# ── ASOSIY IMPORT ─────────────────────────────────────────────────
print("📂 Fayl o'qilmoqda...")
wb = load_workbook(FAYL, read_only=True)
ws = wb.active
rows = list(ws.iter_rows(values_only=True))
print(f"   Jami qatorlar: {len(rows)}")

# Header qatorini tekshirish
if rows and rows[0][0] and str(rows[0][0]).strip().upper() in ('FIO', 'F.I.O', 'ISM', 'TALABA'):
    rows = rows[1:]  # header ni o'tkazib yuborish

print("\n⚙️  Ma'lumotlar tahlil qilinmoqda...")

# Noyob fakultet va guruhlarni yig'ish
fak_cache = {}
guruh_cache = {}
xatolar = []

yangi_fak = yangi_guruh = yangi_talaba = mavjud_talaba = 0

total = len(rows)
print(f"   {total} ta talaba yuklanadi...\n")

for i, row in enumerate(rows, 1):
    try:
        if not row or not row[0]:
            continue

        fio     = str(row[0]).strip() if row[0] else ''
        kurs    = kurs_raqam(row[1]) if len(row) > 1 else 1
        fak_nom = str(row[2]).strip() if len(row) > 2 and row[2] else 'Boshqa'
        grh_nom = str(row[3]).strip() if len(row) > 3 and row[3] else ''

        if not fio or not grh_nom:
            continue

        # Fakultet
        if fak_nom not in fak_cache:
            fak, created = Fakultet.objects.get_or_create(nomi=fak_nom)
            fak_cache[fak_nom] = fak
            if created:
                yangi_fak += 1
        fak = fak_cache[fak_nom]

        # Guruh
        guruh_key = f"{grh_nom}_{fak_nom}"
        if guruh_key not in guruh_cache:
            guruh, created = Guruh.objects.get_or_create(
                nomi=grh_nom,
                defaults={'fakultet': fak, 'kurs': kurs, 'kerakli_sert': 1}
            )
            guruh_cache[guruh_key] = guruh
            if created:
                yangi_guruh += 1
        guruh = guruh_cache[guruh_key]

        # Talaba
        if not Talaba.objects.filter(fio=fio, guruh=guruh).exists():
            Talaba.objects.create(fio=fio, guruh=guruh)
            yangi_talaba += 1
        else:
            mavjud_talaba += 1

        # Progress
        if i % 500 == 0:
            print(f"   {i}/{total} qator ishlandi...")

    except Exception as e:
        xatolar.append(f"{i}-qator: {e}")

# ── NATIJA ────────────────────────────────────────────────────────
print("\n" + "="*50)
print("✅ IMPORT YAKUNLANDI!")
print("="*50)
print(f"  Yangi fakultet  : {yangi_fak} ta")
print(f"  Yangi guruh     : {yangi_guruh} ta")
print(f"  Yangi talaba    : {yangi_talaba} ta")
print(f"  Allaqachon bor  : {mavjud_talaba} ta")
print(f"  Xatolar         : {len(xatolar)} ta")
print("="*50)
print(f"  Jami talaba (bazada): {Talaba.objects.count()} ta")
print(f"  Jami guruh  (bazada): {Guruh.objects.count()} ta")

if xatolar:
    print("\n⚠️  Xatolar:")
    for x in xatolar[:10]:
        print(f"   {x}")
    if len(xatolar) > 10:
        print(f"   ... va yana {len(xatolar)-10} ta xato")
