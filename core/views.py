import csv, io
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required, user_passes_test
from django.contrib.auth.models import User
from django.contrib import messages
from django.http import HttpResponse
from django.db.models import Q, Count, Avg
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from .models import *

def is_admin(u): return u.is_staff or u.is_superuser
def get_profile(u): return getattr(u, 'profile', None)

# ── AUTH ──────────────────────────────────────────────────────────
def login_view(request):
    if request.user.is_authenticated: return redirect('dashboard')
    if request.method == 'POST':
        u = authenticate(request, username=request.POST.get('username',''), password=request.POST.get('password',''))
        if u: login(request, u); return redirect('dashboard')
        messages.error(request, "Login yoki parol noto'g'ri!")
    return render(request, 'core/login.html')

def logout_view(request):
    logout(request); return redirect('login')

# ── DASHBOARD ─────────────────────────────────────────────────────
@login_required
def dashboard(request):
    p = get_profile(request.user)
    now = timezone.now()
    ctx = {
        'jami_sert': Sertifikat.objects.count(),
        'bu_oy': Sertifikat.objects.filter(yaratilgan__year=now.year, yaratilgan__month=now.month).count(),
        'oxirgilar': Sertifikat.objects.select_related('talaba','xodim','oqituvchi')[:8],
        'profile': p,
    }
    if p:
        if p.rol == 'tyutor' and p.guruhlar.exists():
            talabalar = Talaba.objects.filter(guruh__in=p.guruhlar.all())
            jami = talabalar.count()
            yetarli = sum(1 for t in talabalar if t.kpi_foiz() >= 100)
            ctx.update({'jami_talaba': jami, 'yetarli': yetarli, 'foiz': int(yetarli/jami*100) if jami else 0})
        elif p.rol == 'bolim_boshlig' and p.bolim:
            xodimlar = Xodim.objects.filter(bolim=p.bolim)
            jami = xodimlar.count()
            yetarli = sum(1 for x in xodimlar if x.kpi_foiz() >= 100)
            ctx.update({'jami_xodim': jami, 'yetarli': yetarli, 'foiz': int(yetarli/jami*100) if jami else 0})
        elif p.rol == 'kafedra_mudiri' and p.kafedra:
            oqituvchilar = Oqituvchi.objects.filter(kafedra=p.kafedra)
            jami = oqituvchilar.count()
            yetarli = sum(1 for o in oqituvchilar if o.kpi_foiz() >= 100)
            ctx.update({'jami_oqituvchi': jami, 'yetarli': yetarli, 'foiz': int(yetarli/jami*100) if jami else 0})
    return render(request, 'core/dashboard.html', ctx)

# ── RO'YXAT VA SERTIFIKAT YUKLASH ─────────────────────────────────
@login_required
def royxat(request):
    p = get_profile(request.user)
    if not p and not is_admin(request.user):
        messages.error(request, "Sizga bo'lim biriktirilmagan!"); return redirect('dashboard')

    shaxslar = []
    sarlavha = ""

    if p:
        if p.rol == 'tyutor' and p.guruhlar.exists():
            guruhlar = p.guruhlar.all()
            sarlavha = ", ".join(g.nomi for g in guruhlar) + " guruhlari talabalari"
            shaxslar = [{'obj': t, 'tur': 'talaba', 'soni': t.sertifikat_soni(),
                         'kerak': t.guruh.kerakli_sert, 'foiz': t.kpi_foiz()} for t in Talaba.objects.filter(guruh__in=guruhlar).select_related('guruh')]
        elif p.rol == 'bolim_boshlig' and p.bolim:
            sarlavha = f"{p.bolim.nomi} xodimlari"
            shaxslar = [{'obj': x, 'tur': 'xodim', 'soni': x.sertifikat_soni(),
                         'kerak': p.bolim.kerakli_sert, 'foiz': x.kpi_foiz()} for x in Xodim.objects.filter(bolim=p.bolim)]
        elif p.rol == 'kafedra_mudiri' and p.kafedra:
            sarlavha = f"{p.kafedra.nomi} o'qituvchilari"
            shaxslar = [{'obj': o, 'tur': 'oqituvchi', 'soni': o.sertifikat_soni(),
                         'kerak': p.kafedra.kerakli_sert, 'foiz': o.kpi_foiz()} for o in Oqituvchi.objects.filter(kafedra=p.kafedra)]

    return render(request, 'core/royxat.html', {'shaxslar': shaxslar, 'sarlavha': sarlavha, 'profile': p})

@login_required
def sert_yuklash(request, tur, pk):
    p = get_profile(request.user)
    if not p and not is_admin(request.user):
        return redirect('dashboard')

    obj = None
    if tur == 'talaba': obj = get_object_or_404(Talaba, pk=pk)
    elif tur == 'xodim': obj = get_object_or_404(Xodim, pk=pk)
    elif tur == 'oqituvchi': obj = get_object_or_404(Oqituvchi, pk=pk)

    sertlar = Sertifikat.objects.filter(**{tur: obj})

    if request.method == 'POST':
        nomi = request.POST.get('nomi','').strip()
        link = request.POST.get('link','').strip()
        if not nomi or not link:
            messages.error(request, "Nom va link majburiy!")
        else:
            s = Sertifikat(nomi=nomi, link=link,
                tashkilot=request.POST.get('tashkilot',''),
                berilgan=request.POST.get('berilgan') or None,
                yuklagan=request.user)
            setattr(s, tur, obj)
            s.save()
            messages.success(request, "✅ Sertifikat qo'shildi!")
            return redirect('sert_yuklash', tur=tur, pk=pk)

    return render(request, 'core/sert_yuklash.html', {'obj': obj, 'tur': tur, 'sertlar': sertlar})

@login_required
def sert_ochir(request, pk):
    s = get_object_or_404(Sertifikat, pk=pk)
    if s.yuklagan == request.user or is_admin(request.user):
        s.delete(); messages.success(request, "O'chirildi.")
    back = request.META.get('HTTP_REFERER', 'dashboard')
    return redirect(back)

# ── ADMIN: FOYDALANUVCHILAR ───────────────────────────────────────
@login_required
@user_passes_test(is_admin)
def users_list(request):
    users = User.objects.filter(is_staff=False).select_related('profile__guruh','profile__bolim','profile__kafedra').order_by('last_name','first_name')
    return render(request, 'core/users_list.html', {'users': users, 'rol_choices': ROL_CHOICES})

@login_required
@user_passes_test(is_admin)
def user_add(request):
    error = None
    if request.method == 'POST':
        un = request.POST.get('username','').strip()
        pw = request.POST.get('password','').strip()
        fn = request.POST.get('firstname','').strip()
        ln = request.POST.get('lastname','').strip()
        rol = request.POST.get('rol','')
        guruh_id = request.POST.get('guruh')
        bolim_id = request.POST.get('bolim')
        kafedra_id = request.POST.get('kafedra')
        if not all([un,pw,fn,ln,rol]):
            error = "Barcha majburiy maydonlarni to'ldiring!"
        elif User.objects.filter(username=un).exists():
            error = f"'{un}' login allaqachon mavjud!"
        else:
            u = User.objects.create_user(un,'',pw); u.first_name=fn; u.last_name=ln; u.save()
            pr = UserProfile.objects.create(user=u, rol=rol,
                bolim_id=bolim_id or None, kafedra_id=kafedra_id or None)
            guruh_ids = request.POST.getlist('guruhlar')
            if guruh_ids:
                pr.guruhlar.set(guruh_ids)
            messages.success(request, f"✅ '{un}' foydalanuvchisi yaratildi!")
            return redirect('users_list')
    ctx = {'error': error, 'post': request.POST, 'rol_choices': ROL_CHOICES,
           'guruhlar': Guruh.objects.select_related('fakultet').all(),
           'bolimlar': Bolim.objects.all(), 'kafedralar': Kafedra.objects.select_related('fakultet').all()}
    return render(request, 'core/user_add.html', ctx)

@login_required
@user_passes_test(is_admin)
def user_delete(request, pk):
    u = get_object_or_404(User, pk=pk)
    if not u.is_staff: u.delete(); messages.success(request, "O'chirildi.")
    return redirect('users_list')

# ── ADMIN: TUZILMA ────────────────────────────────────────────────
@login_required
@user_passes_test(is_admin)
def tuzilma(request):
    return render(request, 'core/tuzilma.html', {
        'fakultetlar': Fakultet.objects.prefetch_related('kafedralar','guruhlar').all(),
        'bolimlar': Bolim.objects.prefetch_related('xodimlar').all(),
    })

# ── ADMIN: BARCHA SERTIFIKATLAR ───────────────────────────────────
@login_required
@user_passes_test(is_admin)
def admin_sertlar(request):
    qs = Sertifikat.objects.select_related('talaba__guruh','xodim__bolim','oqituvchi__kafedra','yuklagan').all()
    q = request.GET.get('q',''); tur = request.GET.get('tur','')
    if q: qs = qs.filter(Q(talaba__fio__icontains=q)|Q(xodim__fio__icontains=q)|Q(oqituvchi__fio__icontains=q)|Q(nomi__icontains=q))
    if tur == 'talaba': qs = qs.filter(talaba__isnull=False)
    elif tur == 'xodim': qs = qs.filter(xodim__isnull=False)
    elif tur == 'oqituvchi': qs = qs.filter(oqituvchi__isnull=False)
    if request.GET.get('export') == 'csv':
        return export_sert_excel(qs)
    return render(request, 'core/admin_sertlar.html', {'sertlar': qs, 'q': q, 'tur': tur, 'jami': qs.count()})

def export_sert_excel(qs):
    wb = Workbook(); ws = wb.active; ws.title = "Sertifikatlar"
    header = ['#','F.I.O.','Tur',"Bo'lim/Guruh",'Sertifikat nomi','Tashkilot','Berilgan','Link','Yuklagan']
    bold = Font(bold=True); fill = PatternFill("solid", fgColor="1e5fa8")
    for col, h in enumerate(header, 1):
        c = ws.cell(1, col, h); c.font = Font(bold=True, color="FFFFFF"); c.fill = fill; c.alignment = Alignment(horizontal='center')
    for i, s in enumerate(qs, 1):
        ws.append([i, s.egasi_fio(), s.egasi_tur(), s.egasi_bolim(), s.nomi, s.tashkilot,
                   str(s.berilgan or ''), s.link, s.yuklagan.get_full_name() or s.yuklagan.username if s.yuklagan else ''])
    ws.column_dimensions['A'].width = 5; ws.column_dimensions['B'].width = 30
    ws.column_dimensions['C'].width = 12; ws.column_dimensions['D'].width = 30
    ws.column_dimensions['E'].width = 40; ws.column_dimensions['H'].width = 50
    resp = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = 'attachment; filename="sertifikatlar.xlsx"'
    wb.save(resp); return resp

# ── KPI HISOBOT ───────────────────────────────────────────────────
@login_required
def kpi(request):
    p = get_profile(request.user)
    ctx = {'profile': p}
    if is_admin(request.user) or (p and p.rol == 'tyutor'):
        guruh_id = request.GET.get('guruh') or (p.guruhlar.exists() if p and p.rol=='tyutor' else None)
        guruhlar = Guruh.objects.select_related('fakultet').all() if is_admin(request.user) else Guruh.objects.filter(pk=p.guruhlar.exists())
        if guruh_id:
            g = get_object_or_404(Guruh, pk=guruh_id)
            talabalar = [{'obj':t,'soni':t.sertifikat_soni(),'kerak':g.kerakli_sert,'foiz':t.kpi_foiz()} for t in Talaba.objects.filter(guruh=g)]
            ctx.update({'guruh': g, 'talabalar': talabalar, 'guruhlar': guruhlar,
                        'jami': len(talabalar), 'yetarli': sum(1 for t in talabalar if t['foiz']>=100),
                        'o_foiz': int(sum(t['foiz'] for t in talabalar)/len(talabalar)) if talabalar else 0})
        else:
            ctx['guruhlar'] = guruhlar

    if is_admin(request.user) or (p and p.rol == 'bolim_boshlig'):
        bolim_id = request.GET.get('bolim') or (p.bolim_id if p and p.rol=='bolim_boshlig' else None)
        bolimlar = Bolim.objects.all() if is_admin(request.user) else Bolim.objects.filter(pk=p.bolim_id)
        if bolim_id:
            b = get_object_or_404(Bolim, pk=bolim_id)
            xodimlar = [{'obj':x,'soni':x.sertifikat_soni(),'kerak':b.kerakli_sert,'foiz':x.kpi_foiz()} for x in Xodim.objects.filter(bolim=b)]
            ctx.update({'bolim': b, 'xodimlar': xodimlar, 'bolimlar': bolimlar,
                        'jami_x': len(xodimlar), 'yetarli_x': sum(1 for x in xodimlar if x['foiz']>=100),
                        'o_foiz_x': int(sum(x['foiz'] for x in xodimlar)/len(xodimlar)) if xodimlar else 0})
        else:
            ctx.setdefault('bolimlar', bolimlar)

    if is_admin(request.user) or (p and p.rol == 'kafedra_mudiri'):
        kaf_id = request.GET.get('kafedra') or (p.kafedra_id if p and p.rol=='kafedra_mudiri' else None)
        kafedralar = Kafedra.objects.select_related('fakultet').all() if is_admin(request.user) else Kafedra.objects.filter(pk=p.kafedra_id)
        if kaf_id:
            k = get_object_or_404(Kafedra, pk=kaf_id)
            oqituvchilar = [{'obj':o,'soni':o.sertifikat_soni(),'kerak':k.kerakli_sert,'foiz':o.kpi_foiz()} for o in Oqituvchi.objects.filter(kafedra=k)]
            ctx.update({'kafedra': k, 'oqituvchilar': oqituvchilar, 'kafedralar': kafedralar,
                        'jami_o': len(oqituvchilar), 'yetarli_o': sum(1 for o in oqituvchilar if o['foiz']>=100),
                        'o_foiz_o': int(sum(o['foiz'] for o in oqituvchilar)/len(oqituvchilar)) if oqituvchilar else 0})
        else:
            ctx.setdefault('kafedralar', kafedralar)

    if request.GET.get('export') == 'kpi': return export_kpi_excel(ctx)
    return render(request, 'core/kpi.html', ctx)

def export_kpi_excel(ctx):
    wb = Workbook()
    hdr_fill = PatternFill("solid", fgColor="1e5fa8")
    hdr_font = Font(bold=True, color="FFFFFF")
    green_fill = PatternFill("solid", fgColor="27ae60")
    red_fill   = PatternFill("solid", fgColor="e74c3c")
    yellow_fill= PatternFill("solid", fgColor="f39c12")

    def make_sheet(ws, rows, headers):
        for col, h in enumerate(headers, 1):
            c = ws.cell(1, col, h); c.font = hdr_font; c.fill = hdr_fill; c.alignment = Alignment(horizontal='center')
        for i, row in enumerate(rows, 2):
            for col, val in enumerate(row, 1): ws.cell(i, col, val)
            foiz = row[-2]
            color = green_fill if foiz>=100 else yellow_fill if foiz>=50 else red_fill
            ws.cell(i, len(row)-1).fill = color
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(len(str(c.value or '')) for c in col) + 4

    if 'talabalar' in ctx:
        ws = wb.active; ws.title = "Talabalar KPI"
        rows = [(i+1,t['obj'].fio,t['obj'].guruh.nomi,t['soni'],t['kerak'],t['foiz'],f"{t['foiz']}%") for i,t in enumerate(ctx['talabalar'])]
        make_sheet(ws, rows, ['#','F.I.O.','Guruh','Sertifikat soni','Kerak','Foiz (%)','Status'])

    if 'xodimlar' in ctx:
        ws2 = wb.create_sheet("Xodimlar KPI")
        rows = [(i+1,x['obj'].fio,x['obj'].bolim.nomi,x['soni'],x['kerak'],x['foiz'],f"{x['foiz']}%") for i,x in enumerate(ctx['xodimlar'])]
        make_sheet(ws2, rows, ['#','F.I.O.',"Bo'lim",'Sertifikat soni','Kerak','Foiz (%)','Status'])

    if 'oqituvchilar' in ctx:
        ws3 = wb.create_sheet("O'qituvchilar KPI")
        rows = [(i+1,o['obj'].fio,o['obj'].kafedra.nomi,o['soni'],o['kerak'],o['foiz'],f"{o['foiz']}%") for i,o in enumerate(ctx['oqituvchilar'])]
        make_sheet(ws3, rows, ['#','F.I.O.','Kafedra','Sertifikat soni','Kerak','Foiz (%)','Status'])

    resp = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = 'attachment; filename="kpi_hisobot.xlsx"'
    wb.save(resp); return resp

# ── EXCEL IMPORT ──────────────────────────────────────────────────
@login_required
@user_passes_test(is_admin)
def import_page(request):
    ctx = {'guruhlar': Guruh.objects.select_related('fakultet').all(),
           'bolimlar': Bolim.objects.all(),
           'kafedralar': Kafedra.objects.select_related('fakultet').all(),
           'rol_choices': ROL_CHOICES}
    return render(request, 'core/import.html', ctx)

@login_required
@user_passes_test(is_admin)
def import_talabalar(request):
    if request.method != 'POST': return redirect('import_page')
    guruh_id = request.POST.get('guruh')
    f = request.FILES.get('file')
    if not f or not guruh_id:
        messages.error(request, "Fayl va guruhni tanlang!"); return redirect('import_page')
    guruh = get_object_or_404(Guruh, pk=guruh_id)
    try:
        wb = load_workbook(f, read_only=True); ws = wb.active
        count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            fio = str(row[0]).strip() if row[0] else ''
            if not fio or fio.lower() in ('none','—',''): continue
            if not Talaba.objects.filter(fio=fio, guruh=guruh).exists():
                Talaba.objects.create(fio=fio, guruh=guruh); count += 1
        messages.success(request, f"✅ {count} ta talaba qo'shildi ({guruh.nomi})")
    except Exception as e:
        messages.error(request, f"Xatolik: {e}")
    return redirect('import_page')

@login_required
@user_passes_test(is_admin)
def import_xodimlar(request):
    if request.method != 'POST': return redirect('import_page')
    bolim_id = request.POST.get('bolim'); f = request.FILES.get('file')
    if not f or not bolim_id:
        messages.error(request, "Fayl va bo'limni tanlang!"); return redirect('import_page')
    bolim = get_object_or_404(Bolim, pk=bolim_id)
    try:
        wb = load_workbook(f, read_only=True); ws = wb.active; count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            fio = str(row[0]).strip() if row[0] else ''
            lavozim = str(row[1]).strip() if len(row)>1 and row[1] else ''
            if not fio or fio.lower() in ('none','—',''): continue
            if not Xodim.objects.filter(fio=fio, bolim=bolim).exists():
                Xodim.objects.create(fio=fio, bolim=bolim, lavozim=lavozim); count += 1
        messages.success(request, f"✅ {count} ta xodim qo'shildi ({bolim.nomi})")
    except Exception as e: messages.error(request, f"Xatolik: {e}")
    return redirect('import_page')

@login_required
@user_passes_test(is_admin)
def import_oqituvchilar(request):
    if request.method != 'POST': return redirect('import_page')
    kafedra_id = request.POST.get('kafedra'); f = request.FILES.get('file')
    if not f or not kafedra_id:
        messages.error(request, "Fayl va kafedrni tanlang!"); return redirect('import_page')
    kafedra = get_object_or_404(Kafedra, pk=kafedra_id)
    try:
        wb = load_workbook(f, read_only=True); ws = wb.active; count = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            fio = str(row[0]).strip() if row[0] else ''
            lavozim = str(row[1]).strip() if len(row)>1 and row[1] else ''
            if not fio or fio.lower() in ('none','—',''): continue
            if not Oqituvchi.objects.filter(fio=fio, kafedra=kafedra).exists():
                Oqituvchi.objects.create(fio=fio, kafedra=kafedra, lavozim=lavozim); count += 1
        messages.success(request, f"✅ {count} ta o'qituvchi qo'shildi ({kafedra.nomi})")
    except Exception as e: messages.error(request, f"Xatolik: {e}")
    return redirect('import_page')

@login_required
@user_passes_test(is_admin)
def import_users(request):
    if request.method != 'POST': return redirect('import_page')
    f = request.FILES.get('file')
    if not f: messages.error(request, "Fayl tanlang!"); return redirect('import_page')
    try:
        wb = load_workbook(f, read_only=True); ws = wb.active; count = 0; errors = []
        for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), 2):
            try:
                un=str(row[0]).strip(); pw=str(row[1]).strip()
                fn=str(row[2]).strip() if row[2] else ''; ln=str(row[3]).strip() if row[3] else ''
                rol=str(row[4]).strip() if row[4] else ''
                birikma=str(row[5]).strip() if len(row)>5 and row[5] else ''
                if not un or not pw or not rol: continue
                if User.objects.filter(username=un).exists(): errors.append(f"{i}-qator: '{un}' allaqachon bor"); continue
                u=User.objects.create_user(un,'',pw); u.first_name=fn; u.last_name=ln; u.save()
                pr = UserProfile(user=u, rol=rol)
                if rol=='tyutor' and birikma:
                    g=Guruh.objects.filter(nomi__icontains=birikma).first()
                    if g: pr.guruh=g
                elif rol=='bolim_boshlig' and birikma:
                    b=Bolim.objects.filter(nomi__icontains=birikma).first()
                    if b: pr.bolim=b
                elif rol=='kafedra_mudiri' and birikma:
                    k=Kafedra.objects.filter(nomi__icontains=birikma).first()
                    if k: pr.kafedra=k
                pr.save(); count += 1
            except Exception as e: errors.append(f"{i}-qator: {e}")
        msg = f"✅ {count} ta user yaratildi."
        if errors: msg += " ⚠ Xatolar: " + "; ".join(errors[:3])
        messages.success(request, msg)
    except Exception as e: messages.error(request, f"Xatolik: {e}")
    return redirect('import_page')

def shablon_yuklash(request, tur):
    wb = Workbook(); ws = wb.active
    hf = PatternFill("solid", fgColor="1e5fa8"); hfont = Font(bold=True, color="FFFFFF")
    if tur == 'talabalar':
        ws.title = "Talabalar"; headers = ['F.I.O. (majburiy)', 'Izoh (ixtiyoriy)']
        ws.append(['Karimov Akbar Salimovich','']); ws.append(['Rahimova Gulnora Usmonovna',''])
    elif tur == 'xodimlar':
        ws.title = "Xodimlar"; headers = ['F.I.O. (majburiy)', 'Lavozim (ixtiyoriy)']
        ws.append(['Toshmatov Sardor Hamidovich','Bosh mutaxassis']); ws.append(['Yusupova Malika','Hisobchi'])
    elif tur == 'oqituvchilar':
        ws.title = "O'qituvchilar"; headers = ['F.I.O. (majburiy)', 'Lavozim (ixtiyoriy)']
        ws.append(['Mirzayev Jasur Normatovich','Dotsent']); ws.append(['Qodirov Dilshod','Katta o\'qituvchi'])
    elif tur == 'users':
        ws.title = "Users"; headers = ['Login','Parol','Ism','Familiya','Rol','Birikma (guruh/bo\'lim/kafedra nomi)']
        ws.append(['tyutor_mt21','parol123','Akbar','Karimov','tyutor','MT-21-01'])
        ws.append(['bolim_raqamli','parol123','Sardor','Toshmatov','bolim_boshlig',"Raqamli ta'lim markazi"])
        ws.append(['kafedra_mud','parol123','Gulnora','Rahimova','kafedra_mudiri','Ingliz tili kafedrasi'])
    ws.insert_rows(1)
    for col, h in enumerate(headers, 1):
        c = ws.cell(1, col, h); c.font = hfont; c.fill = hf
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(len(str(c.value or '')) for c in col) + 6
    resp = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = f'attachment; filename="shablon_{tur}.xlsx"'
    wb.save(resp); return resp
